from inspect import Attribute
import sqlite3

from string import ascii_letters, ascii_uppercase
from collections import Counter
from os import remove
from unicodedata import name
from xml.etree.ElementInclude import include
from numpy import isin
import openpyxl as xl
from openpyxl.worksheet.table import Table, TableStyleInfo

def RoundToX(num, base):
    return base * round(num / base)
def CalculateMedian(values, medianType = 0.5, sort = True):
    lenValues = len(values)
    midValue = (lenValues + 1) * medianType
    if sort:
        values.sort()
    try:
        if midValue % 1 == 0:
            return values[int(midValue)]
        else:
            midValue = int(midValue)
            return (values[midValue] + values[midValue + 1]) / 2
    except IndexError as e:
        if lenValues <= 2:
            print("not enough values for median")
            return -1
        else:
            return values[int(medianType * 4) - 1]     
def CalculateStandardDeviation(values, mean = False):
    lenValues = len(values)
    if not mean:
        mean = sum(values) / lenValues
    sumSquared = sum([x ** 2 for x in values])

    variation = (sumSquared / lenValues) - mean ** 2
    standardDeviation = variation ** 0.5

    return standardDeviation   
def sign(val):
    if val == 0:
        return 0
    else:
        return 1 if val > 0 else -1


allNodes = {"Match" : ["matchID", "gameID", "replayName", "ballchasingLink", "map", "matchType", 
                "teamSize", "playlistID", "durationCalculated", "durationBallchasing", 
                "overtime", "season", "seasonType", "date", "time", "mmr", "nFrames", "orangeScore", "blueScore", 
                "goalSequence", "neutralPossessionTime", "bTimeGround", "bTimeLowAir", "bTimeHighAir", "bTimeBlueHalf", 
                "bTimeOrangeHalf", "bTimeBlueThird", "bTimeNeutralThird", "bTimeOrangeThird", "bTimeNearWall", "bTimeInCorner", 
                "bTimeOnWall", "bAverageSpeed", "bType", "gameMutatorIndex", "tBlueClumped", "tOrangeClumped", "tBlueIsolated", 
                "tOrangeIsolated", "tBluePossession", "tOrangePossession", "replayTagOne", "replayTagTwo", "replayTagThree", 
                "replayTagFour", "replayTagFive", "startPlayerIndex"],
            "Player" : ["playerID", "matchID", "gameID", "pBallchasingID", "pCalculatedId", "pName", "pPlatform", "pTier", 
                "carName", "titleID", "teamColour", "bUsage", "bPerMinute", "bConsumptionPerMinute", "aAmount", 
                "qCollected", "qStolen", "qCollectedBig", "qCollectedSmall", "qStolenBig", "qStolenSmall", "nCollectedBig", 
                "nCollectedSmall", "nStolenBig", "nStolenSmall", "qOverfill", "qOverfillStolen", "qWasted", "tZeroBoost", "tFullBoost", 
                "tBZeroQuarter", "tBQuaterHalf", "tBHalfUpperQuater", "tBUpperQuaterFull", "aSpeed", "aHitDistance", "aDistanceFromCentre", 
                "dTotal", "tSonicS", "tBoostS", "tSlowS", "tGround", "tLowAir", "tHighAir", "tPowerslide", "nPowerslide", "aPowerslideDuration", 
                "aSpeedPercentage", "aDBall", "aDBallPossession", "aDBallNoPossession", "aDMates", "tDefensiveThird", "tNeutralThird", "tOffensiveThird", 
                "tDefensiveHalf", "tOffensiveHalf", "tBehindBall", "tInFrontBall", "tMostBack", "tMostForward", "goalsConcededLast", "tClosestBall", 
                "tFarthestBall", "tCloseBall", "tNearWall", "tInCorner", "tOnWall", "dHitForward", "dHitBackward", "pTime", "turnovers", "turnoversMyHalf", 
                "turnoversTheirHalf", "wonTurnovers", "aPDuration", "aPHits", "qPossession", "demoInflicted", "demoTaken", "score", "goals", 
                "assists", "saves", "shots", "mvp", "shootingP", "totalHits", "totalPasses", "totalDribbles", "totalDribblesConts", 
                "totalAerials", "totalClears", "isKeyboard", "tBallCam", "qCarries", "qFlicks", "totalCarryT", "totalCarryD", "aCarryT", "totalKickoffs", 
                "numGoBoost", "numnGoFollow", "numGoBall", "numFirstTouch", "aBoostUsed", "fiftyWins", "fiftyLosses", "fiftyDraws", "isBot", "partyLeaderID", 
                "ballchasingStartTime", "ballchasingEndTime", "ballchasingBoostTime", "ballchasingStatTime", "calculatedFirstFrame", "calculatedTimeInGame"]}
retrievalNodes = {"Match" : ["matchID", "gameID", "map", "matchType", 
                "teamSize", "durationCalculated", "durationBallchasing", 
                "overtime", "nFrames", "orangeScore", "blueScore", 
                "goalSequence", "neutralPossessionTime", "bTimeGround", "bTimeLowAir", "bTimeHighAir", "bTimeBlueHalf", 
                "bTimeOrangeHalf", "bTimeBlueThird", "bTimeNeutralThird", "bTimeOrangeThird", "bTimeNearWall", "bTimeInCorner", 
                "bTimeOnWall", "bAverageSpeed", "bType", "gameMutatorIndex", "tBlueClumped", "tOrangeClumped", "tBlueIsolated", 
                "tOrangeIsolated", "tBluePossession", "tOrangePossession", "replayTagOne", "replayTagTwo", "replayTagThree", 
                "replayTagFour", "replayTagFive", "startPlayerIndex"],
                  "Player" : ["playerID", "pBallchasingID", "pCalculatedId", "pName", "pPlatform", "pTier", 
                "carName", "titleID", "teamColour", "bUsage", "bPerMinute", "bConsumptionPerMinute", "aAmount", 
                "qCollected", "qStolen", "qCollectedBig", "qCollectedSmall", "qStolenBig", "qStolenSmall", "nCollectedBig", 
                "nCollectedSmall", "nStolenBig", "nStolenSmall", "qOverfill", "qOverfillStolen", "qWasted", "tZeroBoost", "tFullBoost", 
                "tBZeroQuarter", "tBQuaterHalf", "tBHalfUpperQuater", "tBUpperQuaterFull", "aSpeed", "aHitDistance", "aDistanceFromCentre", 
                "dTotal", "tSonicS", "tBoostS", "tSlowS", "tGround", "tLowAir", "tHighAir", "tPowerslide", "nPowerslide", "aPowerslideDuration", 
                "aSpeedPercentage", "aDBall", "aDBallPossession", "aDBallNoPossession", "aDMates", "tDefensiveThird", "tNeutralThird", "tOffensiveThird", 
                "tDefensiveHalf", "tOffensiveHalf", "tBehindBall", "tInFrontBall", "tMostBack", "tMostForward", "goalsConcededLast", "tClosestBall", 
                "tFarthestBall", "tCloseBall", "tNearWall", "tInCorner", "tOnWall", "dHitForward", "dHitBackward", "pTime", "turnovers", "turnoversMyHalf", 
                "turnoversTheirHalf", "wonTurnovers", "aPDuration", "aPHits", "qPossession", "demoInflicted", "demoTaken", "score", "goals", 
                "assists", "saves", "shots", "mvp", "shootingP", "totalHits", "totalPasses", "totalDribbles", "totalDribblesConts", 
                "totalAerials", "totalClears", "isKeyboard", "tBallCam", "qCarries", "qFlicks", "totalCarryT", "totalCarryD", "aCarryT", "totalKickoffs", 
                "numGoBoost", "numnGoFollow", "numGoBall", "numFirstTouch", "aBoostUsed", "fiftyWins", "fiftyLosses", "fiftyDraws", "isBot", "partyLeaderID", 
                "ballchasingStartTime", "ballchasingEndTime", "ballchasingBoostTime", "ballchasingStatTime", "calculatedFirstFrame", "calculatedTimeInGame", "matchID"]}
analysisNodeDictionary = {"default" : {"analysisType" : 0, "relevancy" : 1},
                          "qOverfillStolen" : {"punishDuplicates" : True},
                          "carName" : {"analysisType" : 1},
                          "scoredFirst" : {"analysisType" : 2},
                          "mvp" : {"analysisType" : 2},
                          
                         }
punishDuplicatesAvg = {"default" : 0.2}
statNodes = {}
percentageAccountValues = {"default" : [0.7, 0.2]}


tagsDictionary = {}
 
class ValueNode:
    def __init__(self, name, percentage = False, calculation = False, teamStat = True, default = -1, valueType = "Player", valueRangeType = 0, tag = None) -> None:
        """Name: Name of Node
           Percentage: Name of Node it is Percentage Of
           Calculation: List, Element 0 is eval string in form "@ + @", where @ are replaced with variables, all other elements are names of variables in calculation
           Team Stat: Whether that stat is relevant to a team/ should be included in the Team Object
           Default: If the fetched value is None, NaN or -1 it is set to this value
           ValueType: What object the node is for
           ValueRangeType: Whether to treat it as a continous curve of values or as a True or False, etc.
           
           Index: Location in retrievalNodes"""
        
        
        self.n = name
        self.p = percentage #name of stat that it is percentage of
        self.c = calculation

        self.tS = teamStat
        self.valueRangeType = valueRangeType
        self.default = default
        self.valueType = valueType
        if name in retrievalNodes[valueType]:
            self.i = retrievalNodes[valueType].index(name)
        else:
            self.i = -1
        
        self.tag = tag
    def __eq__(self, __o: object) -> bool:
        return self.n == __o.n
    def __repr__(s) -> str:
        output = f"Name: {s.n}, Percentage: {s.p}"
        if "rawValue" in s.__dict__:
            output += f", Raw Value: {s.rawValue}, Percentage Of: {s.percentageOf}, Calculated Value: {s.calculatedValue}"
        return output
    def GiveValue(s, rawValue, percentageOf = None, calculationValues = None, individualPlayers = None, teamCalculation = False, gName = None):
        """
        Raw Value: Value of Node
        Percentage Of: What is is percentage of/ will be divided by
        Calculation Values: Stores the values used to calculate the raw values (if needed)
        Individual Players: Stores the individual players behind the sum stat in Team case
        Team Calculation: Lets node know to ignore calculation and other things due to it being teamStat and sum already given
        
        Output:
        
        Returns new node"""
        name = s.n
        if gName:
            name = gName
        node = ValueNode(name, s.p, s.c, s.tS, s.default, s.valueType, s.valueRangeType, s.tag) 
        node.rawValue = rawValue
        node.percentageOf = percentageOf
        node.calculationValues = calculationValues
        node.individualPlayers = individualPlayers


        if node.rawValue in [None, "NaN", -1]:
            node.calculatedValue = s.default
            return node
        cachedValue = node.rawValue
        if node.c and not teamCalculation:
            if node.c == True:
                cachedValue = calculationValues[0]
            else:
                if calculationValues.count(-1) == len(calculationValues):
                    node.calculatedValue = s.default
                    return node
                else:
                    calculationValues = [x if x != -1 else 0 for x in calculationValues]
                    calculationString = node.c[0]
                    for cValue in calculationValues:
                        calculationString = calculationString.replace("@", str(cValue), 1)
                    cachedValue = eval(calculationString)
            node.rawValue = cachedValue
        if node.p and not teamCalculation and calculationValues:
            if percentageOf and -1 not in percentageOf:
                if sum(percentageOf) == 0:
                    cachedValue /= 1
                else:
                    cachedValue /= sum(percentageOf)
            else:
                cachedValue = -1
        node.calculatedValue = cachedValue
        return node
class StatNode:
    def __init__(self, valueNode : ValueNode, values : list, includeValueNode = False) -> None:
        """ValueNode : Example ValueNode of Value
           Values : List of Raw Values (not Value Nodes)"""
        self.name = valueNode.n
        self.valueNode = valueNode
        self.values = [x for x in values if x != -1]
        if includeValueNode:
            self.values.append(valueNode.__dict__[includeValueNode])
        self.numValues = len(self.values)
        if self.numValues > 0:
            self.CalculateStats()
        else:
            print("Zero Values for Stat")
            self.mean, self.mode = -1, -1
    def CalculateStats(self):
        self.values.sort()
        if self.valueNode.valueRangeType == 0:
            self.mean = sum(self.values) / len(self.values)
            self.quartiles = [CalculateMedian(self.values, x, False) for x in (0.25, 0.5, 0.75)]
            self.mode = max(self.values, key = self.values.count)
            self.valuesCounter = Counter(self.values)
            self.standardDeviation = CalculateStandardDeviation(self.values, self.mean)
            try:
                groupValue = statNodes[self.name]["groupValue"]
            except KeyError:
                if self.quartiles[2] - self.quartiles[0] > 5:
                    groupValue = round(0.1 * (self.quartiles[2] - self.quartiles[0]))
                elif self.quartiles[2] - self.quartiles[0] > 1:
                    #groupValue = round(0.1 * (self.__dict__[f"{prefix}Quartiles"][2] - self.__dict__[f"{prefix}Quartiles"][0]), 2)
                    groupValue = round(0.1 * (self.quartiles[2] - self.quartiles[0]), 1)
                else:
                    groupValue = 0.01
                if groupValue == 0:
                    print(self.name)
                    print(self.values)
                    print(self.quartiles)
                    print(0.1 * (self.quartiles[2] - self.quartiles[0]))
                    input()  
            self.groupValue = groupValue
            self.groupedValues = [RoundToX(x, groupValue) for x in self.values]
            self.groupedMode = max(self.groupedValues, key = self.groupedValues.count)
            self.groupedValuesCounter = Counter(self.groupedValues)
        else:
            counterValues = Counter(self.values)
            self.mode = max(counterValues, key = lambda x : counterValues[x])
            self.mean = self.mode
            self.valuesCounter = counterValues
            self.groupedValuesCounter = counterValues
    def OutputValuesStr(s):
        output = "\n"
        
        if s.valueNode.valueRangeType == 0:
            output += f"\nMean : {round(s.mean, 2)}"
            output += f"\nMode : {round(s.mode, 2)}"
            output += f"\nQuartiles: {[round(x, 2) for x in s.quartiles]}"
            output += f"\nStandard Deviation: {round(s.standardDeviation, 2)}"
            output += f"\nGrouped Mode: {round(s.groupedMode, 2)}"
            output += f"\nGroup Value: {round(s.groupValue, 2)}"
        else:
            output += f"\nMean : {s.mean}"
            output += f"\nMode : {s.mode}"
        return output
    def __repr__(s) -> str:
        output = f"{s.name.title()}:"
        output += f"\nAnalysis Type : {s.valueNode.valueRangeType}"
        output += s.OutputValuesStr()
        return output
    def CompareAgainstStatNode(s, otherStatNode):
        """s: Self
        otherStatNode: Stat Node to Compare Against
        
        Output (valueRangeType == 0):
        -Values compared against each other, relative then differential
        Output (valueRangeType != 0):
        -Relative values compared for each count, output in two lists relative and differential"""
        otherStatNode : StatNode
        if s.valueNode.valueRangeType == 0:
            valuesToCompare = ["mean", "quartiles", "standardDeviation", "groupedMode"]
            valuesCompared = {}
            for value in valuesToCompare:
                if value == "quartiles":
                    for i, ourValue in s.__dict__[value]:
                        theirValue = otherStatNode.__dict__[value][i]
                        if theirValue == 0:
                            theirValue = ourValue / -1
                            if theirValue == 0:
                                theirValue = -1
                        valuesCompared[f"{value}{i}"] = [ourValue / theirValue, ourValue - theirValue]    
                else:
                    ourValue = s.__dict__[value]
                    theirValue = otherStatNode.__dict__[value]
                    if theirValue == 0:
                        theirValue = ourValue / -1
                        if theirValue == 0:
                            theirValue = -1
                    valuesCompared[value] = [ourValue / theirValue, ourValue - theirValue]
        else:
            relativeCounts = []
            differentialCount = {}
            relativeDifferenceCount = {}
            for counter in [s.groupedValuesCounter, otherStatNode.groupedValuesCounter]:
                length = sum([x for x in counter.values()])
                relativeCount = {}
                for key in counter:
                    relativeCount[key] = counter[key] / length
                relativeCounts.append(relativeCount)
            while len(relativeCounts[0]) > 0:
                key = relativeCounts[0][relativeCounts[0].keys()[0]]
                ourRelativeValue = relativeCounts[0].pop(key)
                try:
                    theirRelativeValue = relativeCounts[1].pop(key)
                    relativeDifferenceCount[key] = ourRelativeValue / theirRelativeValue
                except KeyError:
                    theirRelativeValue = 0
                    relativeDifferenceCount[key] = -1
                differentialCount[key] = ourRelativeValue - theirRelativeValue
            while len(relativeCounts[1]) > 0:
                key = relativeCounts[1][relativeCounts[1].keys()[0]]
                theirRelativeValue = relativeCounts[1].pop(key)
                relativeDifferenceCount[key] = 0
                differentialCount[key] = -theirRelativeValue
            valuesCompared = {}
            for key in relativeDifferenceCount:
                valuesCompared[key] = [relativeDifferenceCount[key], differentialCount[key]]
        return valuesCompared
            
        
valueNodes = {"Match" : [ValueNode('overtime', valueType = "Match"), 
                    ValueNode('neutralPossessionTime', percentage = "durationCalculated", valueType = "Match"),
                    ValueNode('bTimeGround', percentage = "durationCalculated", valueType = "Match"),
                    ValueNode('bTimeLowAir', percentage = "durationCalculated", valueType = "Match"),
                    ValueNode('bTimeHighAir', percentage = "durationCalculated", valueType = "Match"),
                    ValueNode('bTimeNeutralThird', percentage = "durationCalculated", valueType = "Match"),
                    ValueNode('bTimeNearWall', percentage = "durationCalculated", valueType = "Match"),
                    ValueNode('bTimeInCorner', percentage = "durationCalculated", valueType = "Match"),
                    ValueNode('bTimeOnWall', percentage = "durationCalculated", valueType = "Match"),
                    ValueNode('bAverageSpeed', valueType= "Match"),],
              "Player": [ValueNode('carName', teamStat = False, valueRangeType = 1), 
                    ValueNode('bUsage'),
                    ValueNode('bPerMinute'),
                    ValueNode('bConsumptionPerMinute'),
                    ValueNode('aAmount'),
                    ValueNode('qCollected'),
                    ValueNode('qStolen'),
                    ValueNode('qCollectedBig'),
                    ValueNode('qCollectedSmall'),
                    ValueNode('qStolenBig'),
                    ValueNode('qStolenSmall'),
                    ValueNode('nCollectedBig'),
                    ValueNode('nCollectedSmall'),
                    ValueNode('nStolenBig'),
                    ValueNode('nStolenSmall'),
                    ValueNode('qOverfill'),
                    ValueNode('qOverfillStolen'),
                    ValueNode('qWasted'),
                    ValueNode('tZeroBoost', percentage = "ballchasingBoostTime"),
                    ValueNode('tFullBoost', percentage = "ballchasingBoostTime"),
                    ValueNode('tBZeroQuarter', percentage = "ballchasingBoostTime"),
                    ValueNode('tBQuaterHalf', percentage = "ballchasingBoostTime"),
                    ValueNode('tBHalfUpperQuater', percentage = "ballchasingBoostTime"),
                    ValueNode('tBUpperQuaterFull', percentage = "ballchasingBoostTime"),
                    ValueNode('aSpeed'),
                    ValueNode('aHitDistance'),
                    ValueNode('aDistanceFromCentre'),
                    ValueNode('dTotal'),
                    ValueNode('tSonicS', percentage = "ballchasingStatTime"),
                    ValueNode('tBoostS', percentage = "ballchasingStatTime"),
                    ValueNode('tSlowS', percentage = "ballchasingStatTime"),
                    ValueNode('tGround', percentage = "ballchasingStatTime"),
                    ValueNode('tLowAir', percentage = "ballchasingStatTime"),
                    ValueNode('tHighAir', percentage = "ballchasingStatTime"),
                    ValueNode('tPowerslide', percentage = "ballchasingStatTime"),
                    ValueNode('nPowerslide'),
                    ValueNode('aPowerslideDuration'),
                    ValueNode('aSpeedPercentage'),
                    ValueNode('aDBall'),
                    ValueNode('aDBallPossession'),
                    ValueNode('aDBallNoPossession'),
                    ValueNode('aDMates'),
                    ValueNode('tDefensiveThird', percentage = "ballchasingStatTime"),
                    ValueNode('tNeutralThird', percentage = "ballchasingStatTime"),
                    ValueNode('tOffensiveThird', percentage = "ballchasingStatTime"),
                    ValueNode('tDefensiveHalf', percentage = "ballchasingStatTime"),
                    ValueNode('tOffensiveHalf', percentage = "ballchasingStatTime"),
                    ValueNode('tBehindBall', percentage = "ballchasingStatTime"),
                    ValueNode('tInFrontBall', percentage = "ballchasingStatTime"),
                    ValueNode('tMostBack', percentage = "ballchasingStatTime"),
                    ValueNode('tMostForward', percentage = "ballchasingStatTime"),
                    ValueNode('goalsConcededLast', teamStat = False),
                    ValueNode('tClosestBall', percentage = "ballchasingStatTime"),
                    ValueNode('tFarthestBall', percentage = "ballchasingStatTime"),
                    ValueNode('tCloseBall', percentage = "ballchasingStatTime"),
                    ValueNode('tNearWall', percentage = "ballchasingStatTime"),
                    ValueNode('tInCorner', percentage = "ballchasingStatTime"),
                    ValueNode('tOnWall', percentage = "ballchasingStatTime"),
                    ValueNode('dHitForward'),
                    ValueNode('dHitBackward'),
                    ValueNode('pTime', percentage = "ballchasingStatTime"),
                    ValueNode('turnovers'),
                    ValueNode('turnoversMyHalf'),
                    ValueNode('turnoversTheirHalf'),
                    ValueNode('wonTurnovers'),
                    ValueNode('aPDuration'),
                    ValueNode('aPHits'),
                    ValueNode('qPossession'),
                    ValueNode('demoInflicted'),
                    ValueNode('demoTaken'),
                    ValueNode('score'),
                    ValueNode('goals'),
                    ValueNode('assists'),
                    ValueNode('saves'),
                    ValueNode('shots'),
                    ValueNode('mvp', valueRangeType = 1, teamStat = False),
                    ValueNode('shootingP'),
                    ValueNode('totalHits'),
                    ValueNode('totalPasses'),
                    ValueNode('totalDribbles'),
                    ValueNode('totalDribblesConts'),
                    ValueNode('totalAerials'),
                    ValueNode('totalClears'),
                    ValueNode('tBallCam', percentage = "ballchasingStatTime"),
                    ValueNode('qCarries'),
                    ValueNode('qFlicks'),
                    ValueNode('totalCarryT', percentage = "ballchasingStatTime"),
                    ValueNode('totalCarryD'),
                    ValueNode('aCarryT'),
                    ValueNode('totalKickoffs'),
                    ValueNode('numGoBoost'),
                    ValueNode('numnGoFollow'),
                    ValueNode('numGoBall'),
                    ValueNode('numFirstTouch'),
                    ValueNode('aBoostUsed'),
                    ValueNode('fiftyWins'),
                    ValueNode('fiftyLosses'),
                    ValueNode('fiftyDraws'),
                    ValueNode("fiftyWinRate", percentage = "totalFifties"),
                    ValueNode("fiftyNotLossRate", percentage = "totalFifties", calculation = ["@ + @", "fiftyWins", "fiftyDraws"]),
                    ValueNode('goalParticipation', percentage = "teamGoals", calculation = ["@ + @", "goals", "assists"]),
                    ValueNode('scoredFirst', calculation = True, valueRangeType = 1)],
              "Team": [ValueNode("maxLead"),
                     ValueNode("maxDeficit"),
                     ValueNode("finalLead"),
                     ValueNode("comeback"),
                     ValueNode("choke")]
}
for nodeType in valueNodes:#Match, Player
    nodeDict = {}
    for node in valueNodes[nodeType]:
        nodeDict[node.n] = node
    valueNodes[nodeType] = nodeDict

class AnalysisNode:
    def InitialiseArguments(s, name, kwargs):
        keywordArgs = ["analysisType", "relevancy"] #Analysis Node Arguments
        for kArg in keywordArgs: #For argument in keyword Args
            if kArg in kwargs: #If in given values
                s.__dict__[kArg] = kwargs[kArg] 
            elif name in analysisNodeDictionary and kArg in analysisNodeDictionary[name]: #Checks to see if it's in analysisNodeDict
                s.__dict__[kArg] = analysisNodeDictionary[name][kArg]
            else:
                s.__dict__[kArg] = analysisNodeDictionary["default"][kArg] #Sets to default value
        if False and s.punishDuplicates:
            if name in punishDuplicatesAvg:
                s.punishDuplicates = punishDuplicatesAvg[name]
            else:
                s.punishDuplicates = punishDuplicatesAvg["default"]
        if False and s.percentageAccountForValue:
            if name in percentageAccountValues:
                s.percentageAccountForValue = percentageAccountValues[name]
            else:
                s.percentageAccountForValue = percentageAccountValues["default"]
    def __init__(s, valueNode : ValueNode, againstValues = None, toAnalyse = "calculated", insertValue = True, **kwargs) -> None:
        #Account for Duplicates, Punish Duplicates, Relevancy, 
        s.rBreak = " "
        s.valueNode = valueNode #Setting Value Node 
        s.againstValues = againstValues 
        s.InitialiseArguments(s.valueNode.n, kwargs)
        s.value = valueNode.__dict__[f"{toAnalyse}Value"]
        if s.value == -1: 
            print("invalid value for analysis node")
            return
        if againstValues and againstValues.mean != -1:
            againstValues : StatNode
            if insertValue:
                againstValues.values.append(s.value)
            s.valueIndex = (againstValues.values.index(s.value), len(againstValues.values))
            if s.analysisType == 0:
                try:
                    s.againstAverage = s.value / againstValues.mean
                    s.againstMedian = s.value/ againstValues.quartiles[1]
                    if againstValues.standardDeviation != 0:
                        s.sDAway = (s.value - againstValues.mean) / againstValues.standardDeviation
                    else: 
                        s.sDAway = 0
                    #print(valueNode.n)
                    #print(againstValues.groupedValuesCounter)
                    s.valueRarity = (againstValues.groupedValuesCounter[RoundToX(s.value, againstValues.groupValue)], sum(againstValues.groupedValuesCounter.values()))
                except ZeroDivisionError as e:
                    print(s.analysisType)
                    print(againstValues)
                    raise e
                except AttributeError as e:
                    print(s.valueNode.n)
                    print(againstValues.valueNode.valueRangeType)
                    raise e

            else:
                s.valueRarity = (againstValues.groupedValuesCounter[s.value], sum(againstValues.groupedValuesCounter.values()))
        else:
            print("invalid against nodes")
            return
        s.calculatedRelevancy = s.relevancy * (1.5 - s.valueRarity[0] / s.valueRarity[1])
        if s.analysisType == 0:
            s.calculatedRelevancy *= (1 - s.againstAverage) * (1 - s.againstMedian) * (1 + s.sDAway)
        s.absRelevancy = abs(s.calculatedRelevancy) 
    def __repr__(s) -> str:
        output = f"""{s.valueNode.n}
Index: {s.valueIndex[0]} / {s.valueIndex[1]}
Value Rarity: {s.valueRarity[0]} / {s.valueRarity[1]}"""
        if s.analysisType == 0:
            output += f"""Against Average: {s.againstAverage}
Against Median: {s.againstMedian}
Standard Deviations Away: {s.sDAway}"""
        return output
class HistoricalNode:
    def __init__(s, name, relevancy, value, againstValue, index, analysisType) -> None:
        s.n = name
        s.r = relevancy
        s.v = value
        s.aV = againstValue
        s.i = index
        s.aT = analysisType
        if s.aT == 0:
            s.cR = value / againstValue
            s.cR -= 0.5
            s.cR *= relevancy
        elif s.aT == 1:
            againstValue : Counter
            try:
                s.cR = againstValue[value] / sum(againstValue.values())
            except ZeroDivisionError as e:
                s.cR = 0.5
            
            s.cR -= 0.5
            s.cR *= relevancy
        elif s.aT == 2:
            try:
                s.cR = againstValue[value] / sum(againstValue.values())
            except ZeroDivisionError as e:
                s.cR = 0.5
            s.cR -= 0.5
            s.cR *= relevancy
        else:
            print("cunt")
class Player:
    def __init__(self, playerList, matchList):
        self.pList = playerList
        self.mL = matchList

        self.name = playerList[3]
        self.valueNodes = {}
        
        for node in valueNodes["Player"].values():
            node : ValueNode
            if node.c:
                if node.c == True:
                    if node.n == "scoredFirst":
                        try:
                            calcVariables = [playerList[0] == matchList[11][0]]
                        except (TypeError, IndexError):
                            calcVariables = [-1]
                    else:
                        print(node.n)
                        raise NotImplementedError("Da Fuq 1")
                else:
                    calcVariables = [x.calculatedValue for x in self.valueNodes.values() if x.n in node.c[1:]]
            else:
                calcVariables = None
                rawValue = playerList[node.i]
            if node.p:
                if not isinstance(node.p, list):
                    node.p = [node.p]
                divValues = []
                for divValueOf in node.p:
                    if divValueOf in retrievalNodes["Player"]:
                        divValue = playerList[retrievalNodes["Player"].index(divValueOf)]
                    else:
                        if divValueOf == "teamGoals":
                            teamGoalsIndex = 10 if playerList[8] == "blue" else 9
                            divValue = matchList[teamGoalsIndex]
                        elif divValueOf == "totalFifties":
                            try:
                                divValue = playerList[retrievalNodes["Player"].index("fiftyWins")] + playerList[retrievalNodes["Player"].index("fiftyLosses")] + playerList[retrievalNodes["Player"].index("fiftyDraws")]
                            except TypeError:
                                divValue = -1
                        else:
                            print(node.n)
                            print(node.p)
                            raise NotImplementedError("Da Fuq 2")
                    if divValue in [None, "NaN"]:
                        divValue = -1
                    divValues.append(divValue)
            else:
                divValues = None       
            self.valueNodes[node.n] = node.GiveValue(rawValue, None, calcVariables)
            if node.p:
                self.valueNodes[f"{node.n}%{node.p}"] = node.GiveValue(rawValue, divValues, calcVariables, gName = f"{node.n}%{node.p}")
class PlayerHistoric(Player):
    countForHistoric = 3
    def __init__(s, players, intensiveStats = True):
        s.players = players
        s.statNodes = {}
        s.numAppearances = len(players)
        s.iS = intensiveStats
        s.id = players[0].pList[1]
        for stat in valueNodes["Player"].values():
            statList = [x.valueNodes[stat.n].calculatedValue for x in players]
            s.statNodes[stat.n] = StatNode(stat, statList)
        if intensiveStats:
            s.goalSequencesDictionary = {}
            for playerMatch in players:
                colour = playerMatch.mL[8]
                orangeWin = playerMatch.mL[9] > playerMatch.mL[10]
                win = int(orangeWin if colour == "orange" else not orangeWin)
                try:
                    goalSequence = [int (x) for x in playerMatch.mL[11].split(",")]
                except (TypeError, AttributeError):
                    continue
                except ValueError:
                    print(f"Invalid Goal Sequence Initiation: {playerMatch.mL[11]}")
                    continue
                teamSize = playerMatch.mL[4]
                if len(goalSequence) == 0:
                    print(f"Invalid Goal Sequence Len: {goalSequence}")
                    continue
                if goalSequence[0] > 8:
                    print(f"Goal Sequence IDs out of Bounds: {goalSequence}")
                    continue
                if colour == "orange":
                    goalSequence = [0 if x > teamSize else 1 for x in goalSequence]
                else:
                    goalSequence = [0 if x <= teamSize else 1 for x in goalSequence]
                #0 a goal for you, 1 a goal for them
                currentScore = [0, 0]
                for goal in goalSequence:
                    currentScore[goal] += 1
                    try:
                        s.goalSequencesDictionary[str(currentScore)][win] += 1
                    except KeyError:
                        s.goalSequencesDictionary[str(currentScore)] = [0, 0]
                        s.goalSequencesDictionary[str(currentScore)][win] += 1
                    except TypeError as e:
                        print(s.goalSequencesDictionary)
                        print(s.goalSequencesDictionary[currentScore])
                        print(s.goalSequencesDictionary[currentScore][win])
                        print(currentScore)
                        print(win)
                        raise e
                
                


            return
            allIDs = [x.mL[0] for x in players]
            allPlayers = [x for x in intensiveStats if x.mL[0] in allIDs]

            s.aN = {}

            for stat in PlayerHistoric.analysisNodes:
                if stat.aT != 0:
                    continue
                otherPStats = [x.nodes[stat.n].v for x in allPlayers]
                aAvg = sum(otherPStats) / len(otherPStats)
                s.aN[stat.n] = aAvg
    def __eq__(self, o: object) -> bool:
        if isinstance(o, PlayerHistoric):
            return self.id == o.id
        else:
            return self.id == o
    def GetHistoricPlayer(playerObject : Player, allPlayers):
        playerAppearances = [x for x in allPlayers if x.pList[0] == playerObject.pList[0]]
        if len(playerAppearances) > PlayerHistoric.countForHistoric:
            return PlayerHistoric(playerAppearances)
        return False
class Team:
    def __init__(self, players, matchList) -> None:
        self.players = players
        self.valueNodes = {}
        self.mL = matchList
        self.colour = players[0].pList[8]
        playersValues = [list(x.valueNodes.values()) for x in players]
        for i, node in enumerate(players[0].valueNodes.values()):
            node : ValueNode
            if not node.tS:
                continue
            try:
                playerStatValues = [x[i].calculatedValue for x in playersValues if x[i].calculatedValue != -1]
                if len(playerStatValues) == 0:
                    calculatedValue = -1
                else:
                    calculatedValue = sum(playerStatValues) / len(playerStatValues)
            except TypeError as e:
                if node.n == "scoredFirst":
                    playerStatValues = [x[i].calculatedValue for x in playersValues]
                    calculatedValue = True in playerStatValues
                else:
                    print("Error")
                    print(node.n)
                    print([x[i].v for x in playersValues if x[i].v not in [-1, None]])
                    raise e
            self.valueNodes[node.n] = node.GiveValue(calculatedValue, individualPlayers = playerStatValues, teamCalculation = True)

        teamPlayerIDs = [x.pList[0] for x in players]

        goalSequence = matchList[11]
        if goalSequence:
            goalSequence = goalSequence.split(",")
            currentScores = [0]
            for goal in goalSequence:
                currentScore = currentScores[-1]
                if int(goal) in teamPlayerIDs:
                    currentScore += 1
                else:
                    currentScore -= 1
                currentScores.append(currentScore)

            maxLeadNode = valueNodes["Team"]["maxLead"].GiveValue(max(currentScores))
            minLeadNode = valueNodes["Team"]["maxDeficit"].GiveValue(min(currentScores))
            finalLeadNode = valueNodes["Team"]["finalLead"].GiveValue(currentScores[-1])

            comebackNode = valueNodes["Team"]["comeback"].GiveValue(minLeadNode.calculatedValue if minLeadNode.calculatedValue < 0 and finalLeadNode.calculatedValue > 0 else 0)
            chokeNode = valueNodes["Team"]["choke"].GiveValue(maxLeadNode.calculatedValue if maxLeadNode.calculatedValue > 1 and finalLeadNode.calculatedValue < 0 else 0)
            leadNodes = [maxLeadNode, minLeadNode, finalLeadNode, comebackNode, chokeNode]
            for node in leadNodes:
                self.valueNodes[node.n] = node
class TeamHistoric(Team):
    countForHistoric = 3 
class Match:
    def __init__(self, matchList) -> None:
        self.mL = matchList
        self.valueNodes = {}
        for node in valueNodes["Match"].values():
            rawValue = matchList[node.i]
            if node.p:
                divValue = matchList[retrievalNodes["Match"].index(node.p)]
            else:
                divValue = None
            self.valueNodes[node.n] = node.GiveValue(rawValue, divValue)

    def __repr__(self) -> str:
        output = ""
        for node in self.nodes:
            output += f"{node} : {self.nodes[node]}\n"
        return output
class Cell:
    def __init__(self, x, y, value) -> None:
        self.x = x
        self.y = y
        self.v = value
    def SelfToPosition(self):
        return f"{ascii_uppercase[self.x]}{self.y + 1}"
    def PlaceValue(self, xlSheet):
        xlSheet[self.SelfToPosition()].value = self.v
class ReplayAnalysis:
    def __init__(self, loadReplays = True, args = None):
        #self.dbFile = r"d:\Users\tom\Documents\Visual Studio Code\Python Files\RocketReplayAnalysis\RocketReplayAnalysis\Database\replayDatabase.db"
        self.dbFile = r"D:\Users\tom\Documents\Programming Work\Python\RocketReplayAnalysis\Database\replayDatabase.db"
        self.CreateConnection(self.dbFile)
        self.replays = []
        self.filePath = r"d:\Users\tom\Documents\Visual Studio Code\Python Files\RocketReplayAnalysis\RocketReplayAnalysis\Database\analysisExcelConnection.xlsx"
        #self.filePath = r"D:\Users\tom\Documents\Programming Work\Python\RocketReplayAnalysis\Database\analysisExcelConnection.xlsx"
        self.altdbFile = r"D:\Users\tom\Documents\Programming Work\Python\RocketReplayAnalysis\Database\analysisOutputDatabase.db"
        #self.altdbFile = r"d:\Users\tom\Documents\Visual Studio Code\Python Files\RocketReplayAnalysis\RocketReplayAnalysis\Database\analysisOutputDatabase.db"
        #self.altdbFile = ":memory:"
        self.altConn = False
        if loadReplays:
            self.LoadReplays(args)
    def CreateConnection(self, dbFile):
        print(f"Connected to {dbFile}")
        self.conn = sqlite3.connect(dbFile)
        self.c = self.conn.cursor()
        print(f"SQLite3 Version: {sqlite3.version}")
    def CreateAltConnection(self, dbFile, delete = True):
        print(f"Creating Alternate Connection to {dbFile}")
        if delete:
            remove(dbFile)
            f = open(dbFile, "w")
            f.close()
        self.altConn = sqlite3.connect(dbFile)
        self.altC = self.altConn.cursor()
    def GetReplayDB(self, replayID, getTeams = True):
        if replayID < 0:
            executeSTR = f"SELECT matchID FROM matchTable ORDER BY matchID DESC;"
            self.c.execute(executeSTR)
            replayID *= -1
            replayID = self.c.fetchmany(replayID)[replayID - 1][0]
        executeSTR = f"SELECT {', '.join(retrievalNodes['Match'])} from matchTable WHERE matchID = {replayID}"
        self.c.execute(executeSTR)
        matchDetails = self.c.fetchone()

        executeSTR = f"SELECT {', '.join(retrievalNodes['Player'])} from playerMatchTable WHERE matchID = {replayID}"
        self.c.execute(executeSTR)

        players = [x for x in self.c.fetchall()]
        if getTeams:
            players = [Player(x, matchDetails) for x in players]
            return Match(matchDetails), players, [Team([x for x in players if x.pList[8] == "blue"], matchDetails), Team([x for x in players if x.pList[8] == "orange"], matchDetails)]
        else:
            return Match(matchDetails), [Player(x, matchDetails) for x in players]    
    def GetReplay(self, replayID, getTeams = True):
        match = [x for x in self.matches if x.mL[0] == replayID][0]
        players = [x for x in self.players if x.mL[0] == replayID]
        if getTeams:
            teams = [x for x in self.teams if x.mL[0] == replayID]
            return match, players, teams
        return match, players
    def LoadReplays(self, tagsToLoad, num = -1, loadTeams = True, instantiateHistoricPlayers = True, instantiateHistoricTeams = True, loadStatNodes = True):
        tagsSTR = ""
        if tagsToLoad:
            tagsSTR = f""" WHERE {'and'.join([f'{x[0]} = "{x[1]}"' if isinstance(x[1], str) else f'{x[0]} = {x[1]}' for x in tagsToLoad])}"""
        executeSTR = f"SELECT matchID FROM matchTable ORDER BY matchID DESC{tagsSTR};"

        self.c.execute(executeSTR)


        if num < 1:
            matchIDs = [x[0] for x in self.c.fetchall()]
        else:
            matchIDs = [x[0] for x in self.c.fetchmany(num)]
        matchIDsSTR = f"({', '.join([str(x) for x in matchIDs])})"
        matchNodesToSelectSTR = ", ".join(retrievalNodes["Match"])
        playerNodesToSelectSTR = ", ".join(retrievalNodes["Player"])

        self.c.execute(f"SELECT {matchNodesToSelectSTR} FROM matchTable WHERE matchID in {matchIDsSTR}")
        matchLists = self.c.fetchall()

        self.c.execute(f"SELECT {playerNodesToSelectSTR} FROM playerMatchTable WHERE matchID in {matchIDsSTR}")
        playerLists = self.c.fetchall()
        
        self.matches = []
        self.matches = [Match(x) for x in matchLists]
        self.matchesDict = {x.mL[0] : x for x in self.matches}
        self.players = []
        for player in playerLists:
            self.players.append(Player(player, self.matchesDict[player[-1]].mL))
        self.teamsD = {}
        self.statNodes = []

        if loadTeams:
            for player in self.players:
                try:
                    self.teamsD[player.pList[-1]][player.pList[8]].append(player)
                except KeyError:
                    self.teamsD[player.pList[-1]] = {}
                    self.teamsD[player.pList[-1]][player.pList[8]] = [player]
            self.teamsPlayers = []
            for players in self.teamsD.values():
                try:
                    self.teamsPlayers.append(players["orange"])
                except KeyError:
                    pass
                try:
                    self.teamsPlayers.append(players["blue"])
                except KeyError:
                    pass
            self.teams = []
            for team in self.teamsPlayers:
                self.teams.append(Team(team, team[0].mL))
        if instantiateHistoricPlayers:
            self.historicPlayers = []
            playersDict = {}
            for player in self.players:
                id = player.pList[1]
                if id in playersDict:
                    playersDict[id] += 1
                else:
                    playersDict[id] = 1
            historicDict = {}
            for player in self.players:
                id = player.pList[1]
                if playersDict[id] >= PlayerHistoric.countForHistoric:
                    if id in historicDict:
                        historicDict[id].append(player)
                    else:
                        historicDict[id] = [player]
            for hPlayer in historicDict.values():
                self.historicPlayers.append(PlayerHistoric(hPlayer))
        if instantiateHistoricTeams:
            pass
        if loadStatNodes:
            for valueNode in list(valueNodes["Match"].values()) + list(valueNodes["Player"].values()):
                print(valueNode.n)
                if valueNode.c:
                    continue
                valueNode : ValueNode
                if valueNode.valueType == "Player":
                    allValues = [x.valueNodes[valueNode.n].calculatedValue for x in self.players]
                elif valueNode.valueType == "Match":
                    allValues = [x.valueNodes[valueNode.n].calculatedValue for x in self.matches]
                self.statNodes.append(StatNode(valueNode, allValues))
    def OneAgainstManyAnalysis(s, toAnalyse, analyseAgainst):#, rankBy = False):
        objectValueNodes = toAnalyse.valueNodes.values()
        analysisNodes = []
        againstStatNodes = []
        for valueNode in objectValueNodes:
            nodeName = valueNode.n
            #print(nodeName)
            againstValueNodes = [x.valueNodes[nodeName].calculatedValue for x in analyseAgainst]
            againstStatNode = StatNode(valueNode, againstValueNodes, includeValueNode = "calculatedValue")
            analysisNode = AnalysisNode(valueNode, againstStatNode, insertValue = False)
            againstStatNodes.append(againstStatNode)
            analysisNodes.append(analysisNode)
        #if rankBy:
        #    analysisNodes.sort(key = lambda x : x.__dict__[rankBy] * x.relevancy)
        return analysisNodes, againstStatNodes
    def TwoPlayerHistoricAnalysis(s, pToAnalyse, pToAnalyseAgainst, rankBy = False):
        pass
    def ConvertIndexToPosition(s, position):
        return f"{ascii_uppercase[position[0] - 1]}{position[1]}"
    def OneAgainstManyAnalysisExcel(s, analysisNodes, analysedAgainst, startPosition = (1, 1), sheet = None, override = True):
        valueNColumns = ["n", "calculatedValue", "rawValue", "percentageOf"]
        analysisNColumns = [("valueIndex", 2), "againstAverage", "againstMedian", "sDAway", ("valueRarity", 2), "relevancy", "calculatedRelevancy", "absRelevancy", "rBreak"]
        statNColumns = ["mean", ("quartiles", 3), "mode", "standardDeviation", "groupedMode"]
        allColumnsCombined = valueNColumns + analysisNColumns + statNColumns
        sumAdditionalColumnLength = sum([x[1] - 1 for x in allColumnsCombined if isinstance(x, tuple)])
        #valueNData = [[y.valueNode.__dict__[x] for x in valueNColumns] for y in analysisNodes]
        #analysisNData = [[y.__dict__[x] for x in analysisNColumns] for y in analysisNodes]
        #statNData = [[y.__dict__[x] for x in statNColumns] for y in analysedAgainst]
        combinedData = []
        for i, node in enumerate(analysisNodes):
            #print(node.valueNode.n)
            rowToAdd = []
            for x in valueNColumns:
                try:
                    rowToAdd.append(node.valueNode.__dict__[x])
                except KeyError:
                    rowToAdd.append(-1)
            for x in analysisNColumns:
                key = x
                if isinstance(x, tuple):#-1 multiple times for tuples
                    key = x[0]
                try:
                    rowToAdd.append(node.__dict__[key])
                except KeyError:
                    if isinstance(x, tuple):
                        rowToAdd.extend([-1 for _ in range(x[1])])
                    else:
                        rowToAdd.append(-1)
            for x in statNColumns:
                key = x
                if isinstance(x, tuple):
                    key = x[0]
                try:
                    rowToAdd.append(analysedAgainst[i].__dict__[key])
                except KeyError:
                    if isinstance(x, tuple):
                        rowToAdd.extend([-1 for _ in range(x[1])])
                    else:
                        rowToAdd.append(-1)
            combinedData.append(rowToAdd)
        xlWorkbook = xl.load_workbook(s.filePath)
        if not sheet:
            xlSheet = xlWorkbook.active
        else:
            xlSheet = xlWorkbook[sheet]
        additionalLength = 0
        for i, columnName in enumerate(allColumnsCombined):
            if not isinstance(columnName, tuple):
                xlSheet[s.ConvertIndexToPosition((startPosition[0] + i + additionalLength, startPosition[1]))].value = columnName
            else:
                for _ in range(columnName[1]):
                    xlSheet[s.ConvertIndexToPosition((startPosition[0] + i + additionalLength, startPosition[1]))].value = str(columnName[0]) + str(_)
                    additionalLength += 1
                additionalLength -= 1
        #print(combinedData)
        for y, row in enumerate(combinedData):
            additionalLength = 0
            for x, dataPoint in enumerate(row):
                if isinstance(dataPoint, (list, tuple)):
                    for singlePoint in dataPoint:
                        xlSheet[s.ConvertIndexToPosition((startPosition[0] + x + additionalLength, startPosition[1] + y + 1))].value = singlePoint
                        additionalLength += 1
                    additionalLength -= 1
                else:
                    try:
                        xlSheet[s.ConvertIndexToPosition((startPosition[0] + x + additionalLength, startPosition[1] + y + 1))].value = dataPoint
                    except ValueError:
                        xlSheet[s.ConvertIndexToPosition((startPosition[0] + x + additionalLength, startPosition[1] + y + 1))].value = dataPoint[0]
        overallLength = len(valueNColumns) + len(analysisNColumns) + len(statNColumns) + sumAdditionalColumnLength - 1
        endPosition = (startPosition[0] + overallLength, startPosition[1] + len(analysisNodes))
        table = Table(displayName = "OutputAnalysis", ref = f"{s.ConvertIndexToPosition(startPosition)}:{s.ConvertIndexToPosition(endPosition)}")
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        table.tableStyleInfo = style
        try:
            xlSheet.add_table(table)
        except ValueError as e:
            if override:
                del xlSheet.tables["OutputAnalysis"]
                xlSheet.add_table(table)
            else:
                raise e

        xlWorkbook.save(filePath)
    def OutputPlayerHeadToHead(s, comparedStatNodes, ourPlayer, againstPlayer, startPosition = (1, 1), sheet = None, override = True):
        xlWorkbook = xl.load_workbook(s.filePath)
        if not sheet:
            xlSheet = xlWorkbook.active
        else:
            xlSheet = xlWorkbook[sheet]
        allQuantativeComparisons = [x for x in comparedStatNodes if x.valueRangeType == 0]
        dataToTable = {}
        for comparedNode in allQuantativeComparisons:
            dataToTable[comparedNode.valueNode.n] = comparedNode.comparedAnalysis
        s.GenerateTable(dataToTable)
        #table = s.RecurseTable()
    def GenerateTable(s, dataToTable):
        columns = ["name"] + dataToTable[dataToTable.keys()[0]].keys()
        data = []
        for key, item in dataToTable:
            data = [key] + [x for x in item.values()]
    def RecurseTable(self, key, value, x, y):
        """example dict:
        
        dict = {"key1" : value1,
                "key2" : [value2, value3, value4],
                "key3" : {"valueKey1" : value5,
                          "valueKey2" : [value6, value7]}}
        """
        outputValues = [Cell(x, y, key)]
        y += 1
        if isinstance(value, dict):
            for key, val in value.items():
                x, y, recurseReturn = self.RecurseTable(key, val, x, y)
                outputValues += recurseReturn
        elif isinstance(value, (list, tuple)):
            for val in value:
                outputValues.append(Cell(x, y, val))
                x += 1
            x -= 1
        else:
            outputValues.append(Cell(x, y, value))
        x += 1
        return x, y, outputValues
if __name__ == '__main__':
    replayEngine = ReplayAnalysis(False)
    replayEngine.LoadReplays(None, loadStatNodes = False)
    match, players = replayEngine.GetReplay(123, False)
    output, against = replayEngine.OneAgainstManyAnalysis(players[0], players[1:])
    replayEngine.OutputAnalysisExcel(output, against)

